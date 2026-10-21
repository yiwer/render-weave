"""
生成 Excel 导入的两份样例（票据 13）：price-tags.xlsx 与 route-signs.xlsx。

样例用 openpyxl（MIT）写出，与服务里写空白表的代码无关，用来核对服务读得懂别的程序写的 .xlsx。
单元格按人在 Excel 里会填的样子放：价格是数值，促销截止是日期，首班是时刻，货号是文本格式的数字串，
促销有的是 TRUE/FALSE、有的是"是/否"，列表用顿号、逗号、分号或换行分开，图片有的只写资源 ID。
导入后应得到的数据写在同名的 .json 里，测试按它核对。

只在改样例时运行（需要 Python 3 与 openpyxl 3.1）：
    python -m pip install openpyxl==3.1.5
    python fixtures/excel/make-samples.py
"""
import datetime
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).parent


def header(sheet, names):
    sheet.append(names)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = 'A2'


def price_tags():
    rows = json.loads((HERE / 'price-tags.json').read_text(encoding='utf-8'))['rows']
    book = Workbook()
    sheet = book.active
    sheet.title = '商品价签'
    header(sheet, ['渲染项 ID', '品名', '规格', '单位', '产地', '货号', '零售价', '会员价', '促销', '促销截止', '卖点', '备注'])
    on_sale = [True, '否', None, False, '是', True]
    separators = ['、', ',', '、', '、', '、', '；']
    notes = ['新品', None, '散装', None, None, '进口']
    for index, row in enumerate(rows):
        data = row['data']
        ends = data.get('saleEnds')
        sheet.append([
            row['itemId'], data['productName'], data.get('spec'), data['unit'], data.get('origin'), data['itemNo'],
            data['retailPrice'], data.get('memberPrice'), on_sale[index],
            datetime.date.fromisoformat(ends) if ends else None,
            separators[index].join(data.get('features', [])) or None, notes[index],
        ])
    for cell in sheet['F'][1:]:
        cell.number_format = '@'
    for cell in sheet['J'][1:]:
        cell.number_format = 'yyyy/m/d'
    book.save(HERE / 'price-tags.xlsx')


def route_signs():
    rows = json.loads((HERE / 'route-signs.json').read_text(encoding='utf-8'))['rows']
    book = Workbook()
    main = book.active
    main.title = '线路牌（单线路单方向）'
    header(main, ['线路号', '线路号图片', '票制', '票价（元）', '运营方', '首班', '末班', '运营说明', '班车类型', '班次时刻', '停靠站点'])
    stops = book.create_sheet('停靠站点')
    header(stops, ['停靠站点', '站名', '地标', '所在道路', '站名上方的标识'])
    for index, row in enumerate(rows):
        data = row['data']
        key = data['routeNo']
        first = datetime.time.fromisoformat(data['firstDeparture'])
        image = data.get('routeNoImage')
        main.append([
            data['routeNo'],
            # 第一行只写资源 ID，第二行写完整的 res: 引用
            (image.removeprefix('res:') if index == 0 else image) if image else None,
            data.get('fareType'), data.get('fare'), data.get('operator'), first, data.get('lastDeparture'),
            data.get('serviceNote'), data.get('scheduleType'),
            ('、' if index == 0 else '\n').join(data.get('departureTimes', [])) or None,
            key,
        ])
        main.cell(row=main.max_row, column=6).number_format = 'h:mm' if first.second == 0 else 'h:mm:ss'
        for number, stop in enumerate(data['stops']):
            landmark = stop.get('landmark')
            stops.append([
                key, stop['name'],
                (landmark.removeprefix('res:') if number % 2 else landmark) if landmark else None,
                stop['road'], '、'.join(stop.get('badges', [])) or None,
            ])
    book.save(HERE / 'route-signs.xlsx')


price_tags()
route_signs()
